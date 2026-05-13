from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "paper" / "paper-arxiv.docx"
ARCHITECTURE_FIGURE = ROOT / "paper" / "figures" / "architecture.png"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(8.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, title, headers, rows):
    caption = doc.add_paragraph(title)
    caption.style = "Caption"
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for index, heading in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], heading, bold=True)
        set_cell_shading(table.rows[0].cells[index], "D9EAF7")

    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value)

    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)

    doc.add_paragraph("")


def add_heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def add_para(doc, text):
    paragraph = doc.add_paragraph(text)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.08
    return paragraph


def add_title(doc):
    title = doc.add_paragraph()
    title.style = "Title"
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(
        "MusicBox: Boundary-Aware Rust Integration in Electron for "
        "Performance-Critical Desktop Audio Applications"
    ).bold = True

    author = doc.add_paragraph("Lixi Zhang")
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER

    affiliation = doc.add_paragraph(
        "[Affiliation, city, country, and corresponding author email to be "
        "completed before journal submission]"
    )
    affiliation.alignment = WD_ALIGN_PARAGRAPH.CENTER

    draft = doc.add_paragraph("SCI-oriented journal manuscript draft")
    draft.alignment = WD_ALIGN_PARAGRAPH.CENTER


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    for style_name in ["Normal", "Heading 1", "Heading 2", "Heading 3", "Title", "Caption"]:
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    doc.styles["Normal"].font.size = Pt(10.5)
    doc.styles["Title"].font.size = Pt(16)
    doc.styles["Caption"].font.size = Pt(9)
    doc.styles["Caption"].font.italic = True


def build():
    doc = Document()
    configure_document(doc)
    add_title(doc)

    add_heading(doc, "Abstract")
    add_para(
        doc,
        "Electron makes it practical to build feature-rich desktop audio applications "
        "with web technologies, but performance-critical audio workloads can put "
        "pressure on the Chromium renderer and complicate device-level behavior. "
        "This paper studies MusicBox, an Electron desktop music application whose "
        "interface, library management, lyrics, settings, and plugin-facing workflows "
        "remain in Electron while a Rust native module owns the Windows Audio Session "
        "API (WASAPI) playback data path. The research question is deliberately "
        "narrower than a language-performance comparison: where should the "
        "Electron/Rust boundary be placed so that Electron remains useful for "
        "application engineering while native code handles high-frequency audio work?",
    )
    add_para(
        doc,
        "The evaluation combines a main architecture matrix with supplementary IPC, "
        "format-generalization, seek-stress, and long-stability matrices. The "
        "current protocol deliberately separates playback measurements from the IPC "
        "payload sweep: playback matrices disable payload probing so that pre-playback "
        "IPC allocation and garbage-collection effects do not contaminate CPU, memory, "
        "or lifecycle results. Short playback matrices include warm-up repetitions "
        "that are retained for audit but excluded from condition statistics. The "
        "harness records process-level memory, renderer working set, lifecycle "
        "timings, native render counters, seek command outcomes, logs, quality flags, "
        "audio-file checksums, run metadata, and raw samples. Earlier pilot results "
        "showed lower renderer working set for the Rust/WASAPI path and exposed "
        "seek-induced native underruns; those pilot results motivate the formal "
        "protocol but are not treated as unconditional proof of backend superiority.",
    )
    add_para(
        doc,
        "The results support a boundary-placement contribution: the tested "
        "Rust/WASAPI placement is useful in an Electron audio application when native "
        "code owns the high-frequency data plane, Electron remains on the control "
        "plane, and the system records enough evidence to separate steady playback, "
        "boundary cost, device fallback, and interactive stress behavior. The study "
        "does not claim universal native superiority, cross-device WASAPI generality, "
        "or underrun-free behavior under all interaction patterns.",
    )
    add_para(
        doc,
        "Keywords: Electron; Rust; desktop audio; WASAPI; Node-API; WebAudio; "
        "benchmark; software architecture",
    )

    add_heading(doc, "1. Introduction")
    add_para(
        doc,
        "Desktop audio applications combine two requirements that often pull "
        "architecture in different directions. Users expect responsive interfaces, "
        "media-library workflows, settings, lyrics, plugins, metadata editing, and "
        "operating-system integration. At the same time, playback must preserve "
        "timing, buffer ownership, device negotiation, and predictable behavior "
        "under long-running use. Electron is attractive for the first group of "
        "requirements because it provides Chromium-based rendering, Node.js "
        "integration in the main process, and a mature desktop packaging ecosystem "
        "[1]. However, moving an entire playback data path into the renderer can "
        "increase memory pressure and tie audio behavior to a process that also "
        "owns user-interface responsibilities.",
    )
    add_para(
        doc,
        "A less informative framing would be to compare Rust and browser-side "
        "JavaScript or WebAudio as if the result were mainly about language speed. "
        "Such a comparison would risk confirming an expectation rather than "
        "producing a useful architecture result. MusicBox instead asks whether an "
        "Electron application can retain its web-based productivity while moving "
        "only the performance-critical audio data plane into Rust. The important "
        "design decision is boundary placement: continuous sample movement, device "
        "buffers, and render-loop counters belong to the native backend, while "
        "commands, status, and user workflows remain in Electron.",
    )
    add_para(
        doc,
        "This paper evaluates that decision through a reproducible benchmark harness. "
        "The harness runs controlled backend conditions, stores raw per-run data, "
        "captures console and native logs, computes condition summaries, and "
        "produces quality reports. The evaluation intentionally includes positive "
        "and negative evidence. It reports lower measured memory in the tested native "
        "rows and long-running playback continuity for the tested native backend, "
        "but it also reports device-format fallback messages and seek-induced "
        "underruns. This gives the "
        "work a clearer research "
        "contribution than a simple native-versus-web comparison: it identifies a "
        "useful boundary, quantifies its cost, and defines its present limits.",
    )
    add_para(
        doc,
        "The paper makes five contributions: (i) a boundary-aware Electron/Rust "
        "architecture for desktop audio applications; (ii) a reproducible benchmark "
        "methodology that separates playback telemetry from IPC boundary cost, "
        "records lifecycle latency, process-level memory, renderer working set, "
        "native render counters, seek command behavior, logs, file checksums, and "
        "quality flags; (iii) formal matrices with round-robin execution and warm-up "
        "exclusion; (iv) real-music CC0 fixtures and generated derivatives that make "
        "the evidence auditable; and (v) a constrained interpretation that separates "
        "steady playback, large-payload IPC limits, format generalization, frequent "
        "seek stress, and long-running playback continuity.",
    )

    add_heading(doc, "2. Research Questions and Claims")
    add_para(
        doc,
        "The evaluation is organized around three research questions rather than a "
        "general language-performance comparison.",
    )
    add_para(
        doc,
        "RQ1: Can a Rust/WASAPI data plane reduce renderer and application memory "
        "pressure compared with a WebAudio data path under the same Electron "
        "application shell?",
    )
    add_para(
        doc,
        "RQ2: Is the Electron-to-native boundary cheap enough for control-plane "
        "operations, while still discouraging continuous audio-sample transfer "
        "across IPC?",
    )
    add_para(
        doc,
        "RQ3: Can the benchmark harness expose operational limits, such as "
        "device-format fallback and underruns, instead of hiding them behind "
        "aggregate success or failure metrics?",
    )
    add_para(
        doc,
        "The corresponding claim is intentionally narrow: a Rust-enhanced Electron "
        "architecture is useful for performance-critical desktop audio when Rust "
        "owns the high-frequency audio data path, Electron remains on the "
        "low-frequency control plane, and benchmark evidence includes process "
        "metrics, lifecycle timings, native counters, and logs. The claim is not "
        "that Rust is universally faster than WebAudio, nor that the current backend "
        "is underrun-free under every interaction pattern.",
    )

    add_heading(doc, "3. Related Work and Research Gap")
    add_heading(doc, "3.1 Web-based desktop applications and WebAudio", level=2)
    add_para(
        doc,
        "Electron applications use Chromium renderer processes for user interfaces "
        "and a Node.js-capable main process for desktop integration [1]. WebAudio "
        "provides a high-level audio API based on an audio graph, with processing "
        "primarily performed by the browser implementation and direct script-based "
        "processing also supported [5]. This stack is valuable for interactive "
        "applications because it exposes a portable audio abstraction inside the "
        "same environment as the interface. For a desktop music player, however, "
        "decoded buffers, renderer process memory, and browser-managed audio "
        "services can become part of the performance surface that must be measured "
        "instead of assumed away.",
    )
    add_heading(doc, "3.2 Native audio APIs and operating-system boundaries", level=2)
    add_para(
        doc,
        "WASAPI lets Windows client applications manage the flow of audio data "
        "between an application and an endpoint device, including session control "
        "and rendering endpoint buffers [3]. It is therefore a natural target for a "
        "native data plane on Windows. Its device-dependent behavior is also a "
        "reason to evaluate shared and exclusive modes separately: exclusive-mode "
        "support, format negotiation, fallback decisions, and underruns can depend "
        "on the endpoint and driver. MusicBox treats these behaviors as part of the "
        "evidence rather than as implementation details hidden behind a generic "
        "playback success flag.",
    )
    add_heading(doc, "3.3 Native extension boundary", level=2)
    add_para(
        doc,
        "Node-API provides a stable ABI for building native add-ons and insulating "
        "them from changes in the underlying JavaScript engine [2]. Rust is used for "
        "the native data plane because it supports systems-level implementation with "
        "explicit ownership while still being exposed to Electron through Node-API "
        "[4,10]. The design question is not whether native extensions are possible; it "
        "is what should cross the extension boundary. MusicBox uses Node-API for "
        "control operations, status, and lifecycle calls, while continuous audio "
        "buffers and device timing remain native.",
    )
    add_heading(doc, "3.4 Measurement methodology", level=2)
    add_para(
        doc,
        "Performance evaluation can be misleading when it relies on too few "
        "repetitions, aggregates away outliers, or fails to record the environment "
        "and raw data. Prior systems work has emphasized statistical rigor in "
        "repeated performance evaluation [7] and has shown that measurement setups "
        "can produce wrong conclusions even when no obvious programming error is "
        "present [8]. The MusicBox harness follows that lesson pragmatically: every "
        "run has a manifest, configuration snapshot, raw result file, samples, logs, "
        "condition summaries, quality flags, and excluded-run tables. Electron "
        "process metrics are obtained from app.getAppMetrics(), which reports memory "
        "and CPU statistics for the processes associated with the application [6].",
    )
    add_heading(doc, "3.5 Research gap", level=2)
    add_para(
        doc,
        "The middle ground between a browser-only audio path and a fully native "
        "desktop application is under-documented. Developers need evidence about "
        "where to place the Electron/Rust boundary, how much IPC costs, whether "
        "renderer memory changes materially, and which operational limits remain "
        "after native integration. Cross-platform native audio systems such as "
        "PortAudio show a long-standing need to abstract platform-specific audio "
        "interfaces [11], and comparative work on web-based desktop frameworks has "
        "studied broad resource costs [16]. MusicBox addresses a narrower gap for a "
        "Windows desktop music-player setting: it contributes system-level evidence "
        "for a boundary-aware architecture rather than a general benchmark ranking "
        "of programming languages or audio APIs.",
    )

    add_heading(doc, "4. System Architecture")
    add_para(
        doc,
        "MusicBox is organized into three layers. The renderer layer implements the "
        "visible interface, playback controls, settings, library views, lyrics, and "
        "plugin-facing panels. The Electron main-process layer manages windows, IPC "
        "handlers, dialogs, local services, metadata helpers, and benchmark "
        "orchestration. The Rust native layer implements the audio backend, including "
        "WASAPI initialization, share-mode selection, audio decoding, resampling, "
        "buffer management, render callbacks, seeking, and native render statistics.",
    )
    add_para(
        doc,
        "The central architectural rule is that high-frequency audio data should not "
        "cross the Electron boundary. Electron sends commands and receives status; "
        "Rust owns the buffers and the device clock. This rule turns IPC into a "
        "control-plane mechanism rather than a sample-streaming mechanism. It also "
        "makes the benchmark falsifiable: if IPC payloads become large, the latency "
        "sweep should make that cost visible.",
    )
    add_para(
        doc,
        "The architecture also treats logging as part of the system design. Native "
        "initialization logs device mix format, fallback decisions, output mode, "
        "buffer sizes, underruns, render errors, and lifecycle events. The benchmark "
        "harness parses those logs to distinguish fatal errors, expected warnings, "
        "and environment-specific fallback. Figure 1 summarizes the boundary "
        "placement.",
    )
    if ARCHITECTURE_FIGURE.exists():
        doc.add_picture(str(ARCHITECTURE_FIGURE), width=Inches(5.9))
        caption = doc.add_paragraph(
            "Figure 1. Boundary-aware MusicBox architecture. Electron owns "
            "presentation and control-plane services; Rust/WASAPI owns the playback "
            "data plane and native render counters."
        )
        caption.style = "Caption"

    add_heading(doc, "5. Implementation")
    add_heading(doc, "5.1 Electron layer", level=2)
    add_para(
        doc,
        "The Electron main process is structured around controllers and services. "
        "Controllers register IPC handlers for window management, application "
        "actions, dialogs, WebAudio controls, native audio controls, and benchmark "
        "execution. This keeps renderer code behind the preload boundary and avoids "
        "direct Node.js access from the renderer. For benchmark runs, the application "
        "starts a minimal benchmark page, executes the requested backend condition, "
        "samples process metrics, writes raw results, and closes the run.",
    )
    add_heading(doc, "5.2 Rust native audio engine", level=2)
    add_para(
        doc,
        "The Rust engine is exposed through a Node-API module. It initializes WASAPI, "
        "selects shared or exclusive output mode, negotiates supported device "
        "formats, loads MP3 and FLAC files, performs resampling when source and "
        "output formats differ, writes to render buffers, handles play, pause, stop, "
        "and seek, and exposes counters such as callbacks, frames written, underruns, "
        "render errors, and seek clears. These counters make it possible to "
        "distinguish successful playback with clean rendering from playback that "
        "completed but experienced operational stress.",
    )
    add_heading(doc, "5.3 Benchmark automation", level=2)
    add_para(
        doc,
        "The benchmark automation is controlled by JSON matrices and a one-command "
        "runner. Each run receives a unique label and a timestamped output directory. "
        "Raw output includes configuration, user agent, run metadata, audio-file "
        "fingerprints, lifecycle timings, optional IPC timings, sampled process "
        "snapshots, render statistics, position data, stdout, stderr, and console "
        "logs. The matrix runner records the execution order, build/runtime "
        "fingerprints, and warm-up rows. The summarizer produces run-level and "
        "condition-level CSV files with means, standard deviations, approximate 95% "
        "confidence intervals, working-set slopes, seek command statistics, and log "
        "checks. Warm-up rows remain in the run table but are excluded from condition "
        "statistics. The figure generator produces SVG summaries with confidence "
        "intervals where available, and the quality-report generator records "
        "low-quality sampling flags and fatal log patterns.",
    )

    add_heading(doc, "6. Experimental Evaluation")
    add_heading(doc, "6.1 Measurement principles and environment", level=2)
    add_para(
        doc,
        "The experiments compare architectural placements under the same Electron "
        "shell. WebAudio conditions keep playback in the browser-side audio path. "
        "Native conditions route playback through Rust/WASAPI in shared or exclusive "
        "mode, with the requested share mode passed directly to native initialization. "
        "The IPC-sweep matrix measures the cost of Electron IPC without playback, "
        "while playback matrices set IPC payload iterations to zero. The "
        "redistributable source FLAC is the Open "
        "Goldberg Variations recording released under CC0 1.0; in the local fixture "
        "set it is a 96 kHz stereo file, while the MP3 transcode is 44.1 kHz stereo. "
        "The 48 kHz WAV, 96 kHz FLAC, and long-stability fixtures are generated from "
        "that CC0 music source with FFmpeg [9]. The 30 min long-stability fixture is "
        "a looped real-music derivative, not a synthetic sine signal. The fixture "
        "generator plus manifest document the paths, durations, checksums, and FFmpeg "
        "settings used to recreate them.",
    )
    add_para(
        doc,
        "The benchmark records whole-application working set, renderer working set, "
        "main-process memory, GPU and utility process memory, application CPU, "
        "lifecycle timings, IPC latencies, native counters, seek results, and log "
        "quality. Underrun and render-error counters are native-backend counters; "
        "WebAudio rows are therefore marked as not applicable rather than interpreted "
        "as browser-level underrun telemetry. The tables below report condition means. "
        "Seek latency is an API command-duration metric, not an acoustic output-settling "
        "latency. WebAudio loadTrack reads the full file through the Electron path and "
        "decodes an AudioBuffer; native loadTrack opens/probes the file and streams "
        "decoding during playback. Standard deviations, approximate 95% confidence "
        "intervals, raw samples, and per-run logs remain available in the CSV and "
        "raw-result artifacts.",
    )
    add_table(
        doc,
        "Table 1. Experimental environment.",
        ["Item", "Value"],
        [
            ["Operating system", "Microsoft Windows 11 Home, version 10.0.26200, 64-bit"],
            ["CPU", "12th Gen Intel Core i7-12700H, 14 physical cores, 20 logical processors"],
            ["Memory", "15.7 GB visible memory"],
            [
                "Audio endpoint",
                "Native logs report playback through Speakers (USB Audio Device); "
                "other devices observed included Realtek High Definition Audio and "
                "NVIDIA Virtual Audio Device (WDM)",
            ],
            ["Software versions", "Node v24.15.0; Electron 41.2.1; rustc 1.94.1; cargo 1.94.1"],
            ["Repository revision", "git commit 5ad5a5b at experiment time"],
            ["Primary result date", "Raw batch directories timestamped 2026-05-04"],
        ],
    )

    add_heading(doc, "6.2 Evidence chain", level=2)
    add_para(
        doc,
        "The formal evaluation consists of a main architecture matrix and four "
        "supplementary matrices. A separate sanity matrix checks the end-to-end "
        "benchmark pipeline and is not treated as a paper-result matrix. The evidence "
        "chain below defines the intended role of each batch after rerunning the "
        "revised protocol. Earlier pilot batches are useful for debugging and for "
        "identifying threats to validity, but formal manuscript values should be "
        "taken from batches whose manifests show disabled playback IPC probing, "
        "round-robin order, warm-up handling where configured, audio checksums, and "
        "native build fingerprints.",
    )
    add_table(
        doc,
        "Table 2. Evidence chain used by the manuscript.",
        ["Purpose", "Batch directory", "Runs", "Main outcome"],
        [
            [
                "Main architecture matrix",
                "new formal rerun",
                "54 planned: 9 warm-up + 45 measured",
                "Primary WebAudio versus Rust/WASAPI memory, lifecycle, and native-counter evidence; playback IPC probing disabled.",
            ],
            [
                "IPC payload sweep",
                "new formal rerun or unchanged IPC-only control batch",
                "5",
                "Boundary-only latency from 0 B to 8 MB; defines control-plane limits.",
            ],
            [
                "Format generalization",
                "new formal rerun",
                "36 planned: 9 warm-up + 27 measured",
                "MP3 44.1 kHz, WAV 48 kHz, and FLAC 96 kHz runs under the same round-robin and no-playback-IPC controls.",
            ],
            [
                "Frequent seek robustness",
                "new formal rerun",
                "24 planned: 6 warm-up + 18 measured",
                "Seek command success, seek command duration, and native underrun/render-error counters under periodic interaction.",
            ],
            [
                "Long-running playback continuity",
                "new formal rerun",
                "9",
                "30 min looped real-music FLAC for WebAudio, WASAPI shared, and WASAPI exclusive under matching duration and input.",
            ],
        ],
    )

    add_heading(doc, "6.3 Main architecture matrix", level=2)
    add_para(
        doc,
        "The main matrix is designed to test the central boundary-placement claim "
        "under the least confounded playback setup available in this application. "
        "It uses WebAudio, WASAPI shared, and WASAPI exclusive conditions over the "
        "same audio inputs, disables IPC payload probing during playback, uses "
        "round-robin execution, and excludes configured warm-up runs from condition "
        "statistics. The 48 kHz stereo WAV row is the primary baseline because it "
        "minimizes codec and resampling confounds; compressed and non-48 kHz rows "
        "are interpreted as application-workload rows rather than pure backend "
        "microbenchmarks.",
    )
    add_para(
        doc,
        "Lifecycle timings should be interpreted as application-visible benchmark "
        "path measurements, not as bitwise-equivalent decoder microbenchmarks. "
        "WebAudio loadTrack includes full-file transfer through the Electron path "
        "and AudioBuffer decoding. Native loadTrack measures the native file/probe "
        "path and then streams decoding during playback. Therefore loadTrack and "
        "memory results support claims about the tested MusicBox implementation "
        "strategies, not intrinsic limits of WebAudio or WASAPI.",
    )
    add_table(
        doc,
        "Table 3. Main architecture matrix reporting template.",
        ["Condition", "Measured runs", "Coverage", "App WS MB mean (CI)", "Renderer WS MB mean (CI)", "Load scope", "Native underruns"],
        [
            ["WebAudio WAV 48 kHz", "5", "from formal rerun", "from CSV", "from CSV", "full-file read + AudioBuffer decode", "n/a"],
            ["WASAPI shared WAV 48 kHz", "5", "from formal rerun", "from CSV", "from CSV", "native probe + streaming decode", "from final native stats"],
            ["WASAPI exclusive WAV 48 kHz", "5", "from formal rerun", "from CSV", "from CSV", "native probe + streaming decode", "from final native stats"],
            ["Additional FLAC/MP3 rows", "5 each", "from formal rerun", "from CSV", "from CSV", "application workload", "from final native stats"],
        ],
    )

    add_heading(doc, "6.4 IPC payload sweep", level=2)
    add_para(
        doc,
        "The IPC payload sweep extended the boundary-cost measurement to larger "
        "payloads. Across five boundary-only runs and 300 iterations per payload "
        "size, mean latency was 0.130 ms for 0 B, 0.122 ms for 1 KB, 0.251 ms for "
        "64 KB, 1.877 ms for 1 MB, 6.724 ms for 4 MB, and 12.950 ms for 8 MB. "
        "The small-payload results support Electron IPC for control-plane commands "
        "such as play, stop, seek, and status. The 4 MB and 8 MB results support the "
        "opposite conclusion for sample-like data in this architecture: continuous "
        "large-payload transfer through IPC is not a suitable design path.",
    )

    add_heading(doc, "6.5 Format generalization", level=2)
    add_para(
        doc,
        "The format-generalization matrix repeats the comparison over MP3 44.1 kHz, "
        "WAV 48 kHz, and FLAC 96 kHz inputs. It uses the same no-playback-IPC and "
        "round-robin controls as the main matrix, with warm-up rows excluded from "
        "condition statistics. These rows test whether the observed application-level "
        "memory and lifecycle patterns survive common codec and sample-rate changes. "
        "They do not isolate codec implementation quality, because WebAudio and the "
        "native path use different load/decode strategies.",
    )
    add_table(
        doc,
        "Table 4. Format-generalization reporting template.",
        [
            "Input format",
            "WebAudio renderer MB",
            "WASAPI shared renderer MB",
            "WASAPI exclusive renderer MB",
            "Native underruns",
            "Native render errors",
        ],
        [
            ["WAV 48 kHz", "from formal rerun", "from formal rerun", "from formal rerun", "from final native stats", "from final native stats"],
            ["FLAC 96 kHz", "from formal rerun", "from formal rerun", "from formal rerun", "from final native stats", "from final native stats"],
            ["MP3 44.1 kHz", "from formal rerun", "from formal rerun", "from formal rerun", "from final native stats", "from final native stats"],
        ],
    )

    add_heading(doc, "6.6 Frequent seek robustness", level=2)
    add_para(
        doc,
        "The seek-robustness matrix added 18 runs with one seek every 5 s for "
        "180 s after warm-up. The seek metric is the benchmarked API-call duration, "
        "not acoustic output latency or first-audible-frame latency. A successful "
        "seek event means that the command returned without error; it does not prove "
        "that output buffers had already settled at the new position. This matrix is "
        "therefore used to expose command-path failures, native underruns, render "
        "errors, and sample coverage under interaction stress, rather than to claim "
        "perceptual seek latency.",
    )
    add_table(
        doc,
        "Table 5. Frequent seek stress reporting template.",
        ["Condition", "Runs", "Seek success", "Mean seek API ms", "Underruns/run", "Render errors"],
        [
            ["WebAudio FLAC seek", "3", "from formal rerun", "from formal rerun", "n/a", "n/a"],
            ["WebAudio MP3 seek", "3", "from formal rerun", "from formal rerun", "n/a", "n/a"],
            ["WASAPI shared FLAC seek", "3", "from formal rerun", "from formal rerun", "from final native stats", "from final native stats"],
            ["WASAPI shared MP3 seek", "3", "from formal rerun", "from formal rerun", "from final native stats", "from final native stats"],
            ["WASAPI exclusive FLAC seek", "3", "from formal rerun", "from formal rerun", "from final native stats", "from final native stats"],
            ["WASAPI exclusive MP3 seek", "3", "from formal rerun", "from formal rerun", "from final native stats", "from final native stats"],
        ],
    )

    add_heading(doc, "6.7 Long-running playback continuity", level=2)
    add_para(
        doc,
        "The long-stability matrix uses the same 30 min looped real-music FLAC "
        "fixture for WebAudio, WASAPI shared, and WASAPI exclusive conditions. This "
        "replaces the earlier synthetic-sine and unequal-duration pilot design. The "
        "purpose is to test playback continuity, sample coverage, application working "
        "set, renderer working set, memory slope, native underruns, and native render "
        "errors over a longer continuous workload. Because WebAudio holds a decoded "
        "AudioBuffer while the native path streams through the Rust decoder, memory "
        "differences in this matrix should be interpreted as current implementation "
        "strategy differences rather than as a universal property of browser or native "
        "audio APIs.",
    )
    add_table(
        doc,
        "Table 6. Long-running playback-continuity reporting template.",
        [
            "Condition",
            "Runs",
            "Duration",
            "App WS MB",
            "Renderer WS MB",
            "App slope MB/min",
            "Underruns / errors",
        ],
        [
            ["WebAudio looped music FLAC", "3", "30 min", "from formal rerun", "from formal rerun", "from formal rerun", "n/a / n/a"],
            ["WASAPI shared looped music FLAC", "3", "30 min", "from formal rerun", "from formal rerun", "from formal rerun", "from final native stats"],
            ["WASAPI exclusive looped music FLAC", "3", "30 min", "from formal rerun", "from formal rerun", "from formal rerun", "from final native stats"],
        ],
    )

    add_heading(doc, "6.8 Data quality and auditability", level=2)
    add_para(
        doc,
        "A formal batch is treated as publication-grade only if all planned measured "
        "runs are present, warm-up rows are excluded from condition summaries, no "
        "included run has renderer errors, fatal log patterns are zero or explicitly "
        "justified, playback runs have adequate sample coverage, and all native "
        "underruns or render errors are reported directly. The CSV exporter writes "
        "physical CRLF line breaks and UTF-8 with a byte-order mark so that paths and "
        "log excerpts open correctly in common spreadsheet applications. "
        "Memory-delta fields report last-minus-first net changes, so negative values "
        "can occur when a process ends a run below its initial sampled memory level. "
        "The manifest records audio checksums, runtime details, git status, and the "
        "native module fingerprint. These controls are necessary because pilot "
        "batches exposed issues that can otherwise lead to overinterpretation: "
        "malformed CSV rows, ambiguous IPC timing, unequal long-duration workloads, "
        "and startup mode-switch ambiguity.",
    )

    add_heading(doc, "7. Discussion")
    add_para(
        doc,
        "The architectural interpretation is deliberately narrower than a backend "
        "speed ranking. The contribution is not that Rust is generally faster than "
        "WebAudio. The contribution is that MusicBox defines and measures a specific "
        "boundary: Electron owns interface and control-plane complexity, while the "
        "native module owns the playback data plane and device-facing counters. A "
        "formal result supports this contribution only to the extent that the measured "
        "rows preserve the revised controls: playback IPC disabled, warm-up excluded, "
        "round-robin order, matching long-duration inputs, and explicit reporting of "
        "native underruns and render errors.",
    )
    add_para(
        doc,
        "The IPC payload sweep clarifies why this boundary matters. Small control "
        "messages were low-latency in this harness, but 4 MB and 8 MB payloads cost "
        "materially more than control messages in the pilot sweep. This is acceptable "
        "as a diagnostic measurement, but it would be a poor basis for continuous "
        "audio-sample movement. The revised protocol keeps this measurement separate "
        "from playback rows so that IPC allocation and garbage collection do not bias "
        "memory or CPU summaries. The architecture is therefore not a generic native "
        "acceleration wrapper; it is a data-placement rule.",
    )
    add_para(
        doc,
        "The stability interpretation also remains bounded. Long playback can support "
        "a continuity claim only for the tested build, machine, audio endpoint, input "
        "fixture, and duration. Frequent seek stress is a separate interaction test: "
        "seek command success does not imply click-free or underrun-free acoustic "
        "output. Any underruns observed there should be treated as negative evidence "
        "about buffering and seek-clear behavior, not averaged away as harmless noise.",
    )
    add_para(
        doc,
        "Log auditing remains central. Without logs and counters, device-format "
        "fallback messages and seek-induced underruns could be misread as hidden "
        "failures or "
        "ignored entirely. The benchmark turns those behaviors into inspectable "
        "evidence. This auditability distinguishes the study from a simple "
        "native-versus-web timing demonstration.",
    )

    add_heading(doc, "8. Threats to Validity")
    add_para(
        doc,
        "The evaluation was performed on one Windows machine and one audio-device "
        "environment. WASAPI behavior, especially exclusive-mode format support and "
        "underrun behavior, can vary across devices and drivers. The added "
        "format-generalization and long-stability matrices reduce the original "
        "evidence weakness, but they do not prove cross-device or cross-OS generality.",
    )
    add_para(
        doc,
        "The long-stability matrix strengthens playback-duration evidence but remains "
        "bounded. The revised matrix uses a matching 30 min looped real-music fixture "
        "for WebAudio, WASAPI shared, and WASAPI exclusive conditions, but it still "
        "covers one machine, one operating-system family, one primary audio endpoint, "
        "and one long-duration musical source. Memory slopes should not be used as "
        "proof of leak-free behavior. Longer tests, more varied real-world libraries, "
        "and multi-device validation would be needed for a stronger deployment-level "
        "claim.",
    )
    add_para(
        doc,
        "The benchmark mode intentionally uses a minimal page, which reduces unrelated "
        "UI noise but may underrepresent worst-case interaction with heavy renderer "
        "views, plugins, lyrics rendering, or media-library scans. The frequent-seek "
        "stress partially addresses interaction behavior, but it does not cover all "
        "user workflows. The paper therefore claims architectural usefulness, not a "
        "complete end-user performance guarantee.",
    )
    add_para(
        doc,
        "The experiments use a legally redistributable CC0 music source and generated "
        "derivatives for specific formats and durations. This improves reproducibility, "
        "but a single source recording cannot cover all decoder and metadata paths "
        "encountered in arbitrary music collections. Future work should add a legally "
        "redistributable corpus with varied codec settings, bitrates, tags, cover art, "
        "and sample-rate combinations. The current benchmark also does not capture "
        "the acoustic output signal; it therefore cannot prove audio fidelity, channel "
        "correctness, or first-audible-frame latency without a separate capture or "
        "signal-comparison experiment.",
    )

    add_heading(doc, "9. Code and Data Availability")
    add_para(
        doc,
        "The MusicBox source code, benchmark scripts, matrix configurations, generated "
        "fixtures, raw result files, condition summaries, figures, log-check reports, "
        "and quality reports should be released with a public repository tag or "
        "archival DOI before journal submission. The current local evidence package "
        "is organized under paper/experiments/runs, with one timestamped directory "
        "per batch. Each batch contains manifest.json, config.snapshot.json, raw run "
        "folders, benchmark-runs.csv, benchmark-conditions.csv, benchmark-log-check.csv, "
        "figures, and quality-report.md. This release requirement follows general "
        "research-software practice for making code, data, dependencies, and the "
        "specific cited software version inspectable [12-15].",
    )
    add_para(
        doc,
        "The source audio is documented by the Open Goldberg Variations "
        "Wikimedia Commons file page and released under CC0 1.0. The derived and "
        "looped audio fixtures are documented by "
        "scripts/benchmarks/audio/fixture-manifest.json and can be recreated with "
        "scripts/benchmarks/generate-benchmark-fixtures.js. A final "
        "submission should include the exact commit hash, operating-system details, "
        "dependency versions, fixture-generation command, and all matrix commands. If "
        "the target journal "
        "requires a replication package, the benchmark runner and raw results should "
        "be archived together rather than summarized only in the manuscript.",
    )

    add_heading(doc, "10. Conclusion")
    add_para(
        doc,
        "This paper presented MusicBox, a boundary-aware Electron/Rust architecture "
        "for desktop audio applications. The central contribution is not a general "
        "claim that Rust is faster than WebAudio. The contribution is an architecture "
        "and evidence package showing that Electron can remain responsible for "
        "application and interface complexity while Rust/WASAPI owns the playback "
        "data plane.",
    )
    add_para(
        doc,
        "The revised experimental design makes the evidence more defensible than the "
        "pilot protocol: playback rows no longer include pre-playback IPC payload "
        "probing, short matrices include auditable warm-up rows that are excluded from "
        "condition means, Native initialization uses the requested WASAPI share mode "
        "directly, and the long-stability matrix uses matching 30 min looped real-music "
        "inputs across WebAudio and native modes. These changes are necessary before "
        "formal SCI claims are made from the benchmark results.",
    )
    add_para(
        doc,
        "The defensible conclusion after formal rerun should remain bounded: the "
        "architecture is meaningful if the measured MusicBox implementation shows "
        "lower renderer pressure or better continuity under the stated workloads, but "
        "the results do not establish universal native superiority, browser-level "
        "underrun telemetry, acoustic seek latency, audio fidelity, or cross-device "
        "WASAPI generality. Interactive seek buffering, output-signal validation, and "
        "multi-device replication remain required before making stronger claims.",
    )

    add_heading(doc, "Acknowledgements")
    add_para(
        doc,
        "No external funding or conflicts of interest are declared in this draft. "
        "This statement should be revised before submission if funding, institutional "
        "support, author contributions, data-availability requirements, or conflicts "
        "of interest apply.",
    )

    add_heading(doc, "References")
    references = [
        "[1] Electron. Electron documentation. https://www.electronjs.org/docs/latest/.",
        "[2] Node.js. Node-API documentation. https://nodejs.org/api/n-api.html.",
        "[3] Microsoft Learn. Windows Audio Session API (WASAPI). "
        "https://learn.microsoft.com/en-us/windows/win32/coreaudio/wasapi.",
        "[4] The Rust Programming Language. https://www.rust-lang.org/.",
        "[5] W3C. Web Audio API. https://www.w3.org/TR/webaudio/.",
        "[6] Electron. app.getAppMetrics API. "
        "https://www.electronjs.org/docs/latest/api/app#appgetappmetrics.",
        "[7] A. Georges, D. Buytaert, and L. Eeckhout. Statistically rigorous Java "
        "performance evaluation. Proceedings of the 22nd Annual ACM SIGPLAN "
        "Conference on Object-Oriented Programming Systems, Languages and "
        "Applications (OOPSLA 2007), 57-76. doi:10.1145/1297027.1297033.",
        "[8] T. Mytkowicz, A. Diwan, M. Hauswirth, and P. F. Sweeney. Producing "
        "wrong data without doing anything obviously wrong. Proceedings of the 14th "
        "International Conference on Architectural Support for Programming "
        "Languages and Operating Systems (ASPLOS XIV), 2009, 265-276. "
        "doi:10.1145/1508244.1508275.",
        "[9] FFmpeg. FFmpeg documentation. https://ffmpeg.org/documentation.html.",
        "[10] N. D. Matsakis and F. S. Klock II. The Rust language. Proceedings of "
        "the 2014 ACM SIGAda Annual Conference on High Integrity Language "
        "Technology (HILT 2014), 103-104. doi:10.1145/2663171.2663188.",
        "[11] R. Bencina and P. Burk. PortAudio - an open source cross platform "
        "audio API. Proceedings of the International Computer Music Conference "
        "(ICMC 2001), 263-266. http://hdl.handle.net/2027/spo.bbp2372.2001.036.",
        "[12] G. Wilson, D. A. Aruliah, C. T. Brown, N. P. Chue Hong, M. Davis, "
        "R. T. Guy, et al. Best practices for scientific computing. PLOS Biology "
        "12(1): e1001745, 2014. doi:10.1371/journal.pbio.1001745.",
        "[13] G. Wilson, J. Bryan, K. Cranston, J. Kitzes, L. Nederbragt, and "
        "T. K. Teal. Good enough practices in scientific computing. PLOS "
        "Computational Biology 13(6): e1005510, 2017. "
        "doi:10.1371/journal.pcbi.1005510.",
        "[14] A. M. Smith, D. S. Katz, K. E. Niemeyer, and FORCE11 Software "
        "Citation Working Group. Software citation principles. PeerJ Computer "
        "Science 2: e86, 2016. doi:10.7717/peerj-cs.86.",
        "[15] R. C. Jimenez, M. Kuzak, M. Alhamdoosh, M. Barker, B. Batut, "
        "M. Borg, et al. Four simple recommendations to encourage best practices "
        "in research software [version 1; peer review: 3 approved]. F1000Research "
        "6:876, 2017. doi:10.12688/f1000research.11407.1.",
        "[16] M. Holowinski and B. Panczyk. Comparative analysis of the technology "
        "used to create multi-platform applications on the example of NW.js and "
        "Electron. Journal of Computer Sciences Institute 17:396-400, 2020. "
        "doi:10.35784/jcsi.2380.",
    ]
    for reference in references:
        add_para(doc, reference)

    doc.core_properties.title = (
        "MusicBox: Boundary-Aware Rust Integration in Electron for "
        "Performance-Critical Desktop Audio Applications"
    )
    doc.core_properties.author = "Lixi Zhang"
    doc.core_properties.subject = "SCI-oriented MusicBox manuscript draft with supplementary experiments"
    doc.core_properties.keywords = "Electron, Rust, WASAPI, WebAudio, benchmark, desktop audio"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
